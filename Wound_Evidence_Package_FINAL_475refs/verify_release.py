#!/usr/bin/env python3
"""Independent CSV-to-matrix verification for final manuscript-aligned 11 descriptive figures."""
from pathlib import Path
from collections import Counter
import csv, json, hashlib
ROOT=Path(__file__).resolve().parent
TABLE={1:1,2:1,3:2,4:2,5:2,6:3,7:3,8:4,9:4,10:5,11:5}
N={1:28,2:28,3:20,4:20,5:20,6:37,7:41,8:45,9:45,10:34,11:34}
ORIGINAL={1:1,2:2,3:4,4:5,5:6,6:7,7:8,8:9,9:10,10:12,11:14}
def csvdict(path):
 with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def csvmat(path):
 with path.open(encoding='utf-8-sig',newline='') as f:return list(csv.reader(f))
def norm(x):return ' '.join(str(x or '').casefold().split())
refs=csvdict(ROOT/'Current_Main_475_References.csv')
assert len(refs)==475
known={int(r['Reference_number']) for r in refs}
assert known==set(range(1,476))
si={}
for r in csvdict(ROOT/'Current_SI_Table_References.csv'):
 si.setdefault(int(r['SI_Table'].lstrip('S')),set()).add(int(r['Reference_number']))
assert [len(si[k]) for k in range(1,6)]==[28,20,41,45,34]
checks=0;status=[]
for num in range(1,12):
 coded=csvdict(ROOT/f'coding/S{num:02d}_study_coding.csv')
 m=csvmat(ROOT/f'matrices/S{num:02d}_matrix.csv')
 refs_in_code={int(r['Main_ref']) for r in coded}
 assert len(refs_in_code)==N[num],(num,'unexpected study count')
 assert refs_in_code<=known and refs_in_code<=si[TABLE[num]]
 assert all(r['Code_0_1'].strip() in ('','0','1') for r in coded)
 counts=Counter()
 for r in coded:counts[norm(r['Group']),norm(r['Domain'])]+=int(r['Code_0_1'] or 0)
 old=ORIGINAL[num];this=0
 for row in m[1:]:
  if old in (6,9,12,14):
   col='Group' if old==9 else 'Domain'
   calc=sum(int(r['Code_0_1'] or 0) for r in coded if norm(r[col])==norm(row[0]))
   assert calc==int(row[1]),(num,row,calc);this+=1
  else:
   for c,value in enumerate(row[1:-1]):
    label=m[0][c+1]
    calc=counts[norm(row[0]),norm(label)]
    assert calc==int(value),(num,row[0],label,calc,value);this+=1
 for ext in ['png','svg']:
  f=ROOT/f'figures/Figure_S{num:02d}_REBUILT.{ext}'
  assert f.exists() and f.stat().st_size>2000,(num,ext)
 checks+=this
 status.append({'figure':f'S{num}','SI_source_table':f'S{TABLE[num]}','unique_references':len(refs_in_code),'cells_checked':this,'status':'PASS'})
assert checks==339,checks
try:
 from openpyxl import load_workbook
 wb=load_workbook(ROOT/'Wound_SI_Journal_Source_Data.xlsx',read_only=True,data_only=True)
 rows=list(wb['09_Recalculation_Check'].values)[1:]
 assert len(rows)==339,len(rows)
 assert all(r[5]=='PASS' and abs(float(r[3])-float(r[4]))<1e-5 for r in rows)
 assert set(f'Matrix_S{k:02d}' for k in range(1,12))<=set(wb.sheetnames)
 wb.close()
except ImportError:pass
report={'status':'PASS','figure_count':11,'matrix_values_independently_recomputed':checks,'source_library_count':475,'figure_checks':status,'important_limit':'Internal numeric/source-cohort concordance only; not independent validation of all original-study classifications. Original S3, S11 and S13 ordinal figures removed.'}
(ROOT/'VALIDATION_REPORT.json').write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
print('PASS:',checks,'numeric checks; 11 figures; 475 numbered registered references; 5 SI source cohorts')
